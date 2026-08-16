from io import BytesIO

from openpyxl import Workbook

from backend.app import create_app
from backend.db import get_db
from backend.services import deduplicate_mobilepay_transactions, find_lineup_player, sync_squad
from backend.models import FineCharge, FineRule, Match, MatchParticipant, MobilePayTransaction, Player, Squad


def make_app(tmp_path):
    class TestConfig:
        SECRET_KEY = "test-secret"
        DATABASE_URL = f"sqlite:///{tmp_path / 'test.db'}"
        CORS_ORIGIN = "http://localhost:5173"
        MAX_CONTENT_LENGTH = 10 * 1024 * 1024
        SESSION_COOKIE_SECURE = False
        SESSION_COOKIE_HTTPONLY = True
        SESSION_COOKIE_SAMESITE = "Lax"
        AUTO_CREATE_SCHEMA = True

    return create_app(TestConfig)


def setup_payload():
    return {
        "clubName": "ØB Klubben",
        "squadName": "Serie 1",
        "dbuClubName": "Øster Sundby Boldklub",
        "dbuSeasonUrl": "https://www.dbu.dk/resultater/hold/4964_485586/kampprogram",
        "seasonName": "2026",
        "seasonStart": "2026-01-01",
        "username": "owner",
        "password": "secret123",
    }


def csrf(response):
    return response.get_json()["csrfToken"]


def test_auth_fines_and_public_dashboard(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    assert owner.get("/api/v1/setup/status").get_json()["needsSetup"] is True
    setup = owner.post("/api/v1/setup", json=setup_payload())
    assert setup.status_code == 201
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]

    player_one = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Player One", "mobilePayName": "Player One"},
        headers={"X-CSRF-Token": token},
    )
    player_two = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Player Two", "mobilePayName": "Player Two"},
        headers={"X-CSRF-Token": token},
    )
    assert player_one.status_code == 201
    assert player_two.status_code == 201
    first_id = player_one.get_json()["player"]["id"]
    second_id = player_two.get_json()["player"]["id"]

    rule = owner.post(
        f"/api/v1/squads/{squad_id}/fine-rules",
        json={"name": "Manglende sko", "description": "Husk udstyret", "amount": 50},
        headers={"X-CSRF-Token": token},
    )
    assert rule.status_code == 201
    rule_id = rule.get_json()["rule"]["id"]

    member = app.test_client()
    registration = member.post(
        "/api/v1/auth/register",
        json={"username": "player", "password": "player123", "playerId": first_id},
    )
    assert registration.status_code == 201
    member_token = csrf(member.get("/api/v1/auth/me"))
    request_response = member.post(
        f"/api/v1/squads/{squad_id}/fine-requests",
        json={"ruleId": str(rule_id), "playerIds": [second_id]},
        headers={"X-CSRF-Token": member_token},
    )
    assert request_response.status_code == 201
    request_id = request_response.get_json()["request"]["id"]

    approved = owner.post(
        f"/api/v1/squads/{squad_id}/fine-requests/{request_id}/approve",
        json={},
        headers={"X-CSRF-Token": token},
    )
    assert approved.status_code == 200
    assert approved.get_json()["request"]["status"] == "approved"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Name", "Type", "Number", "Message", "Amount", "Currency", "Transaction type"])
    sheet.append(["05/06/2026 20:21", "Player One", "Shop", "17018", "", 50, "DKK", "Pay in"])
    sheet.append(["05/06/2026 20:21", "Club shop", "Shop", "17018", "", -250, "DKK", "Pay out"])
    file_data = BytesIO()
    workbook.save(file_data)
    imported = owner.post(
        f"/api/v1/squads/{squad_id}/mobilepay-imports",
        data={"file": (BytesIO(file_data.getvalue()), "mobilepay.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert imported.status_code == 201
    assert imported.get_json()["report"]["created"] == 2

    public = app.test_client().get("/api/v1/public/squads/serie-1/dashboard")
    assert public.status_code == 200
    payload = public.get_json()
    assert payload["players"][0]["totalPaid"] == 50
    assert payload["players"][1]["totalFines"] == 50
    assert payload["expenses"][0]["amount"] == -250
    assert payload["balanceSummary"]["boxBalance"] == -200
    assert payload["balanceSummary"]["ifEveryonePays"] == -150

    transactions = owner.get(f"/api/v1/squads/{squad_id}/transactions").get_json()["transactions"]
    expense = next(item for item in transactions if item["transactionType"] == "pay_out")
    assert expense["amount"] == -250
    removed = owner.delete(
        f"/api/v1/squads/{squad_id}/transactions/{expense['id']}",
        headers={"X-CSRF-Token": token},
    )
    assert removed.status_code == 200
    after_removal = app.test_client().get("/api/v1/public/squads/serie-1/dashboard").get_json()
    assert after_removal["expenses"] == []
    assert after_removal["balanceSummary"]["boxBalance"] == 50

    saved_dates = owner.patch(
        f"/api/v1/squads/{squad_id}/seasons/{payload['season']['id']}",
        json={"startDate": "2025-12-01", "endDate": "2026-12-31"},
        headers={"X-CSRF-Token": token},
    )
    assert saved_dates.status_code == 200
    assert saved_dates.get_json()["season"]["startDate"] == "2025-12-01"
    in_period = owner.get(
        f"/api/v1/squads/{squad_id}/transactions?seasonId={payload['season']['id']}"
    ).get_json()["transactions"]
    assert len(in_period) == 1
    owner.patch(
        f"/api/v1/squads/{squad_id}/seasons/{payload['season']['id']}",
        json={"startDate": "2026-06-06", "endDate": "2026-12-31"},
        headers={"X-CSRF-Token": token},
    )
    outside_period = owner.get(
        f"/api/v1/squads/{squad_id}/transactions?seasonId={payload['season']['id']}"
    ).get_json()["transactions"]
    assert outside_period == []
    outside_dashboard = app.test_client().get("/api/v1/public/squads/serie-1/dashboard").get_json()
    assert outside_dashboard["expenses"] == []
    assert outside_dashboard["balanceSummary"]["boxBalance"] == 0


def test_guest_cannot_mutate(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    response = client.post("/api/v1/setup", json=setup_payload())
    squad_id = response.get_json()["squad"]["id"]
    guest = app.test_client()
    denied = guest.post(f"/api/v1/squads/{squad_id}/charges", json={"playerIds": [1], "title": "Nope", "amount": 1})
    assert denied.status_code == 403


def test_admin_can_manage_fine_rules_and_issued_fines(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    season_id = setup.get_json()["squad"]["currentSeason"]["id"]
    player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Fine Player"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    created_rule = owner.post(
        f"/api/v1/squads/{squad_id}/fine-rules",
        json={"name": "Gamle sko", "description": "Gammel beskrivelse", "amount": 20},
        headers={"X-CSRF-Token": token},
    )
    assert created_rule.status_code == 201
    rule_id = created_rule.get_json()["rule"]["id"]

    updated_rule = owner.patch(
        f"/api/v1/squads/{squad_id}/fine-rules/{rule_id}",
        json={"name": "Glemte sko", "description": "Husk dine støvler", "amount": 35},
        headers={"X-CSRF-Token": token},
    )
    assert updated_rule.status_code == 200
    assert updated_rule.get_json()["rule"]["name"] == "Glemte sko"
    assert updated_rule.get_json()["rule"]["description"] == "Husk dine støvler"
    assert updated_rule.get_json()["rule"]["amount"] == 35

    issued = owner.post(
        f"/api/v1/squads/{squad_id}/fine-requests",
        json={"seasonId": season_id, "ruleId": rule_id, "playerIds": [player["id"]]},
        headers={"X-CSRF-Token": token},
    )
    assert issued.status_code == 201
    assert issued.get_json()["request"]["description"] == "Husk dine støvler"
    dashboard = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()
    charge = dashboard["players"][0]["fines"][0]

    updated_charge = owner.patch(
        f"/api/v1/squads/{squad_id}/charges/{charge['id']}",
        json={"title": "Glemte begge sko", "description": "Rettet af admin", "amount": 10},
        headers={"X-CSRF-Token": token},
    )
    assert updated_charge.status_code == 200
    assert updated_charge.get_json()["charge"]["title"] == "Glemte begge sko"
    assert updated_charge.get_json()["charge"]["description"] == "Rettet af admin"
    assert updated_charge.get_json()["charge"]["amount"] == 10

    deleted_rule = owner.delete(
        f"/api/v1/squads/{squad_id}/fine-rules/{rule_id}",
        headers={"X-CSRF-Token": token},
    )
    assert deleted_rule.status_code == 200
    after_rule_deletion = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()
    assert after_rule_deletion["rules"] == []
    assert after_rule_deletion["players"][0]["fines"][0]["amount"] == 10

    deleted_charge = owner.delete(
        f"/api/v1/squads/{squad_id}/charges/{charge['id']}",
        headers={"X-CSRF-Token": token},
    )
    assert deleted_charge.status_code == 200
    after_charge_deletion = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()
    assert after_charge_deletion["players"][0]["fines"] == []
    assert after_charge_deletion["players"][0]["totalFines"] == 0


def test_owner_can_add_a_second_squad(tmp_path):
    app = make_app(tmp_path)
    client = app.test_client()
    setup = client.post("/api/v1/setup", json=setup_payload())
    response = client.post(
        "/api/v1/squads",
        json={"name": "Serie 2", "dbuClubName": "Øster Sundby", "dbuSeasonUrl": ""},
        headers={"X-CSRF-Token": csrf(setup)},
    )
    assert response.status_code == 201
    assert response.get_json()["squad"]["name"] == "Serie 2"
    assert len(client.get("/api/v1/public/squads").get_json()["squads"]) == 2


def test_owner_can_delete_player_account_without_deleting_financial_data(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Account Player", "mobilePayName": "Account Player"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    charge = owner.post(
        f"/api/v1/squads/{squad_id}/charges",
        json={"playerId": player["id"], "title": "Bevares", "amount": 25},
        headers={"X-CSRF-Token": token},
    )
    assert charge.status_code == 201
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Name", "Type", "Number", "Message", "Amount", "Currency", "Transaction type"])
    sheet.append(["05/06/2026 20:21", "Account Player", "Shop", "123", "", 10, "DKK", "Pay in"])
    file_data = BytesIO()
    workbook.save(file_data)
    imported = owner.post(
        f"/api/v1/squads/{squad_id}/mobilepay-imports",
        data={"file": (BytesIO(file_data.getvalue()), "account-player.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert imported.status_code == 201
    member = app.test_client()
    registered = member.post(
        "/api/v1/auth/register",
        json={"username": "forgotten", "password": "", "playerId": player["id"]},
    )
    assert registered.status_code == 201
    before = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()["players"][0]
    assert before["hasAccount"] is True
    assert before["totalFines"] == 25
    assert before["totalPaid"] == 10

    deleted = owner.delete(
        f"/api/v1/squads/{squad_id}/players/{player['id']}/account",
        headers={"X-CSRF-Token": token},
    )
    assert deleted.status_code == 200
    assert member.get("/api/v1/auth/me").status_code == 401
    after = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()["players"][0]
    assert after["hasAccount"] is False
    assert after["totalFines"] == 25
    assert after["totalPaid"] == 10
    replacement = app.test_client().post(
        "/api/v1/auth/register",
        json={"username": "forgotten", "password": "new", "playerId": player["id"]},
    )
    assert replacement.status_code == 201
    owner_id = setup.get_json()["user"]["id"]
    protected_owner = owner.delete(
        f"/api/v1/users/{owner_id}",
        headers={"X-CSRF-Token": token},
    )
    assert protected_owner.status_code == 409


def test_empty_case_insensitive_password_and_username(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    assert setup.status_code == 201
    member = app.test_client()
    registered = member.post(
        "/api/v1/auth/register",
        json={"username": "SimpleUser", "password": ""},
    )
    assert registered.status_code == 201
    member.post("/api/v1/auth/logout")
    logged_in = member.post(
        "/api/v1/auth/login",
        json={"username": "SIMPLEUSER", "password": ""},
    )
    assert logged_in.status_code == 200
    second = app.test_client()
    assert second.post(
        "/api/v1/auth/register",
        json={"username": "MixedPassword", "password": "MiXeD"},
    ).status_code == 201
    assert second.post(
        "/api/v1/auth/login",
        json={"username": "mixedpassword", "password": "mixed"},
    ).status_code == 200


def test_payment_settings_and_multiple_dbu_sources(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad = setup.get_json()["squad"]
    season_id = squad["currentSeason"]["id"]
    payment = owner.put(
        f"/api/v1/squads/{squad['id']}/payment-settings",
        json={
            "boxNumber": "1783qn",
            "paymentUrl": "https://qr.mobilepay.dk/box/3f19c727-6610-4b18-900a-79e619db34fe/pay-in",
        },
        headers={"X-CSRF-Token": token},
    )
    assert payment.status_code == 200
    source = owner.post(
        f"/api/v1/squads/{squad['id']}/seasons/{season_id}/dbu-sources",
        json={"label": "Andet seniorhold", "url": "https://www.dbu.dk/resultater/hold/9999_999999/kampprogram"},
        headers={"X-CSRF-Token": token},
    )
    assert source.status_code == 201
    dashboard = owner.get(f"/api/v1/squads/{squad['id']}/dashboard").get_json()
    assert dashboard["payment"]["boxNumber"] == "1783QN"
    assert len(dashboard["dbuSources"]) == 2


def test_mobilepay_rows_can_be_rematched_without_reupload(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad = setup.get_json()["squad"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Name", "Type", "Number", "Message", "Amount", "Currency", "Transaction type"])
    sheet.append(["05/05/2026 20:21", "Late MobilePay Name", "Transfer", "123", "", 75, "DKK", "Pay in"])
    file_data = BytesIO()
    workbook.save(file_data)
    imported = owner.post(
        f"/api/v1/squads/{squad['id']}/mobilepay-imports",
        data={"file": (BytesIO(file_data.getvalue()), "late-player.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert imported.status_code == 201
    assert imported.get_json()["report"]["unmatched"] == 1

    player = owner.post(
        f"/api/v1/squads/{squad['id']}/players",
        json={"dbuName": "Short DBU Name", "mobilePayName": "Late MobilePay Name"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    rematched = owner.post(
        f"/api/v1/squads/{squad['id']}/transactions/rematch",
        json={},
        headers={"X-CSRF-Token": token},
    )
    assert rematched.status_code == 200
    assert rematched.get_json()["report"] == {"checked": 1, "matched": 1, "ambiguous": 0, "unmatched": 0}
    transactions = owner.get(
        f"/api/v1/squads/{squad['id']}/transactions?seasonId={squad['currentSeason']['id']}"
    ).get_json()["transactions"]
    assert transactions[0]["allocatedPlayerId"] == player["id"]
    second = owner.post(
        f"/api/v1/squads/{squad['id']}/transactions/rematch",
        json={},
        headers={"X-CSRF-Token": token},
    )
    assert second.get_json()["report"]["checked"] == 0


def test_mobilepay_numeric_identifiers_do_not_create_duplicates(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad = setup.get_json()["squad"]

    def workbook_with_number(number):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Date", "Name", "Type", "Number", "Message", "Amount", "Currency", "Transaction type"])
        sheet.append(["05/05/2026 20:21", "Same Payment", "Transfer", number, "", 75, "DKK", "Pay in"])
        file_data = BytesIO()
        workbook.save(file_data)
        return file_data.getvalue()

    first = owner.post(
        f"/api/v1/squads/{squad['id']}/mobilepay-imports",
        data={"file": (BytesIO(workbook_with_number("4512345678.0")), "first.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    second = owner.post(
        f"/api/v1/squads/{squad['id']}/mobilepay-imports",
        data={"file": (BytesIO(workbook_with_number("4512345678")), "second.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert first.get_json()["report"]["created"] == 1
    assert second.get_json()["report"]["created"] == 0
    assert second.get_json()["report"]["skipped"] == 1

    with app.app_context():
        database_session = get_db()
        original = database_session.query(MobilePayTransaction).one()
        database_session.add(
            MobilePayTransaction(
                squad_id=original.squad_id,
                import_id=original.import_id,
                fingerprint="f" * 64,
                transaction_date=original.transaction_date,
                name=original.name,
                transaction_type_name=original.transaction_type_name,
                number="4512345678.0",
                message=original.message,
                amount_cents=original.amount_cents,
                currency=original.currency,
                transaction_type=original.transaction_type,
                allocation_status=original.allocation_status,
            )
        )
        database_session.commit()
        assert deduplicate_mobilepay_transactions(database_session) == 1
        database_session.commit()
        remaining = database_session.query(MobilePayTransaction).one()
        assert remaining.number == "4512345678"


def test_approver_fine_is_applied_immediately(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Instant Player"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    rule = owner.post(
        f"/api/v1/squads/{squad_id}/fine-rules",
        json={"name": "Sko", "amount": 25},
        headers={"X-CSRF-Token": token},
    ).get_json()["rule"]
    response = owner.post(
        f"/api/v1/squads/{squad_id}/fine-requests",
        json={"ruleId": str(rule["id"]), "playerIds": [player["id"]]},
        headers={"X-CSRF-Token": token},
    )
    assert response.status_code == 201
    assert response.get_json()["request"]["status"] == "approved"
    custom = owner.post(
        f"/api/v1/squads/{squad_id}/fine-requests",
        json={"title": "Anden bøde", "description": "Frit skrevet", "amount": 17, "playerIds": [player["id"]]},
        headers={"X-CSRF-Token": token},
    )
    assert custom.status_code == 201
    custom_request_id = custom.get_json()["request"]["id"]
    assert custom.get_json()["request"]["status"] == "approved"
    dashboard = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()
    assert dashboard["players"][0]["totalFines"] == 42


def test_dbu_lineup_matches_an_existing_abbreviated_name():
    player = Player(dbu_name="Mads Rantzau", mobilepay_name="Mads Rantzau")

    assert find_lineup_player([player], "Mads Rantzau Nielsen") is player


def test_dbu_sync_creates_idempotent_match_charges(tmp_path, monkeypatch):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json={**setup_payload(), "dbuClubName": "Club"})
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    player_ids = []
    for name in ("P1", "P2"):
        response = owner.post(
            f"/api/v1/squads/{squad_id}/players",
            json={"dbuName": name},
            headers={"X-CSRF-Token": token},
        )
        assert response.status_code == 201
        player_ids.append(response.get_json()["player"]["id"])
    for rule_type, amount in (("WIN_FINE", 20), ("SCORED_GOAL", 5), ("CONCEDED_GOAL", 2)):
        response = owner.post(
            f"/api/v1/squads/{squad_id}/fine-rules",
            json={"name": rule_type, "type": rule_type, "amount": amount},
            headers={"X-CSRF-Token": token},
        )
        assert response.status_code == 201

    class FakeDBUClient:
        fetch_calls = 0
        home_score = 2
        away_score = 1

        def __init__(self, **_kwargs):
            pass

        def list_season_match_ids(self, _season_url):
            return ["match-1"]

        def fetch_match(self, _match_id, _club_name):
            from datetime import datetime

            FakeDBUClient.fetch_calls += 1
            return {
                "dbu_id": "match-1",
                "home_club": "Club",
                "away_club": "Visitors",
                "home_score": FakeDBUClient.home_score,
                "away_score": FakeDBUClient.away_score,
                "match_date": datetime(2026, 6, 5, 20, 21),
                "lineup": ["P1", "P2", "P3 Auto", "P4 Auto"],
            }

    monkeypatch.setattr("backend.services.DBUClient", FakeDBUClient)
    with app.app_context():
        database_session = get_db()
        squad = database_session.get(Squad, squad_id)
        first_report = sync_squad(database_session, squad)
        database_session.commit()
        rules = database_session.query(FineRule).count()
        match = database_session.query(Match).one()
        match_id = match.id
        player_count = database_session.query(Player).count()
        unmatched = database_session.query(MatchParticipant).filter(
            MatchParticipant.status == "unmatched"
        ).order_by(MatchParticipant.source_name).all()
        initial_charges = database_session.query(FineCharge).filter(FineCharge.match_id == match_id).order_by(FineCharge.player_id).all()
    assert first_report["createdMatches"] == 1
    assert sum(item["chargesCreated"] for season in first_report["seasons"] for item in season["matches"]) == 2
    assert player_count == 2
    assert [participant.source_name for participant in unmatched] == ["P3 Auto", "P4 Auto"]
    assert rules == 3
    assert match.status == "needs_review"
    assert [(charge.player_id, charge.amount_cents) for charge in initial_charges] == [
        (player_ids[0], 3200),
        (player_ids[1], 3200),
    ]

    with app.app_context():
        database_session = get_db()
        retry_report = sync_squad(database_session, database_session.get(Squad, squad_id))
        database_session.commit()
        retried_charges = database_session.query(FineCharge).filter(FineCharge.match_id == match_id).all()
    assert retry_report["updatedMatches"] == 1
    assert sum(item["chargesCreated"] for season in retry_report["seasons"] for item in season["matches"]) == 0
    assert sum(item["chargesUpdated"] for season in retry_report["seasons"] for item in season["matches"]) == 0
    assert len(retried_charges) == 2
    assert FakeDBUClient.fetch_calls == 2

    locked = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": [player_ids[0]]},
        headers={"X-CSRF-Token": token},
    )
    assert locked.status_code == 200
    restored_dbu_lineup = owner.delete(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        headers={"X-CSRF-Token": token},
    )
    assert restored_dbu_lineup.status_code == 200
    with app.app_context():
        database_session = get_db()
        assert database_session.get(Match, match_id).status == "needs_review"
        assert database_session.query(FineCharge).filter(FineCharge.match_id == match_id).count() == 2

    created_player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "P3 Auto"},
        headers={"X-CSRF-Token": token},
    )
    assert created_player.status_code == 201
    with app.app_context():
        database_session = get_db()
        assert database_session.query(FineCharge).filter(FineCharge.match_id == match_id).count() == 3
        assert database_session.get(Match, match_id).status == "needs_review"

    created_last_player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "P4 Auto"},
        headers={"X-CSRF-Token": token},
    )
    assert created_last_player.status_code == 201
    with app.app_context():
        database_session = get_db()
        assert database_session.query(FineCharge).filter(FineCharge.match_id == match_id).count() == 4
        squad = database_session.get(Squad, squad_id)
        completed_report = sync_squad(database_session, squad)
        database_session.commit()
    assert sum(item["chargesCreated"] for season in completed_report["seasons"] for item in season["matches"]) == 0
    assert completed_report["skippedMatches"] == 1
    assert FakeDBUClient.fetch_calls == 2

    corrected = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": [player_ids[0]]},
        headers={"X-CSRF-Token": token},
    )
    assert corrected.status_code == 200
    assert corrected.get_json()["report"] == {"created": 0, "updated": 0, "removed": 3}
    expanded = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": player_ids},
        headers={"X-CSRF-Token": token},
    )
    assert expanded.status_code == 200
    assert expanded.get_json()["report"] == {"created": 1, "updated": 0, "removed": 0}
    switched = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": [player_ids[1]]},
        headers={"X-CSRF-Token": token},
    )
    assert switched.status_code == 200
    assert switched.get_json()["report"] == {"created": 0, "updated": 0, "removed": 1}
    restored = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": [player_ids[0]]},
        headers={"X-CSRF-Token": token},
    )
    assert restored.status_code == 200
    assert restored.get_json()["report"] == {"created": 1, "updated": 0, "removed": 1}
    washer = owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/washer",
        json={"playerId": player_ids[1]},
        headers={"X-CSRF-Token": token},
    )
    assert washer.status_code == 200
    with app.app_context():
        database_session = get_db()
        assert database_session.query(FineCharge).filter(FineCharge.match_id == match_id).count() == 1
    dashboard = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()
    assert dashboard["matches"][0]["lineupLocked"] is True
    assert dashboard["matches"][0]["participants"][0]["playerName"] == "P1"
    assert dashboard["matches"][0]["washerName"] == "P2"

    incremental = owner.post(
        f"/api/v1/squads/{squad_id}/sync",
        json={"seasonId": setup.get_json()["squad"]["currentSeason"]["id"]},
        headers={"X-CSRF-Token": token},
    )
    assert incremental.status_code == 200
    assert incremental.get_json()["report"]["skippedMatches"] == 1
    assert FakeDBUClient.fetch_calls == 2
    still_locked = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()["matches"][0]
    assert still_locked["lineupLocked"] is True
    assert [item["playerName"] for item in still_locked["participants"]] == ["P1"]

    reset = owner.post(
        f"/api/v1/squads/{squad_id}/sync/reset",
        json={"seasonId": setup.get_json()["squad"]["currentSeason"]["id"]},
        headers={"X-CSRF-Token": token},
    )
    assert reset.status_code == 200
    assert reset.get_json()["report"]["deletedMatches"] == 1
    assert reset.get_json()["report"]["createdMatches"] == 1
    assert FakeDBUClient.fetch_calls == 3
    rebuilt = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()["matches"][0]
    rebuilt_match_id = rebuilt["id"]
    assert rebuilt["lineupLocked"] is False
    assert rebuilt["washerId"] is None
    assert len(rebuilt["participants"]) == 4

    FakeDBUClient.home_score = None
    FakeDBUClient.away_score = None
    with app.app_context():
        database_session = get_db()
        no_result_report = sync_squad(database_session, database_session.get(Squad, squad_id), full=True)
        database_session.commit()
        assert database_session.get(Match, rebuilt_match_id).status == "needs_review"
        assert database_session.query(FineCharge).filter(FineCharge.source == "match").count() == 0
    assert sum(item["chargesRemoved"] for season in no_result_report["seasons"] for item in season["matches"]) == 4


def test_match_charges_include_result_and_goal_rules_for_every_player(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json={**setup_payload(), "dbuClubName": "Club"})
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    season_id = setup.get_json()["squad"]["currentSeason"]["id"]
    player_ids = []
    for name in ("P1", "P2"):
        response = owner.post(
            f"/api/v1/squads/{squad_id}/players",
            json={"dbuName": name},
            headers={"X-CSRF-Token": token},
        )
        player_ids.append(response.get_json()["player"]["id"])

    for rule_type, amount in (
        ("WIN_FINE", 10),
        ("DRAW_FINE", 20),
        ("LOSE_FINE", 30),
        ("SCORED_GOAL", 2),
        ("CONCEDED_GOAL", 5),
    ):
        response = owner.post(
            f"/api/v1/squads/{squad_id}/fine-rules",
            json={"name": rule_type, "type": rule_type, "amount": amount},
            headers={"X-CSRF-Token": token},
        )
        assert response.status_code == 201

    scenarios = (
        ("Club", "Visitors", 3, 1, 2100),
        ("Visitors", "Club", 2, 2, 3400),
        ("Club", "Visitors", 1, 4, 5200),
    )
    for index, (home_club, away_club, home_score, away_score, expected_cents) in enumerate(scenarios, start=1):
        response = owner.post(
            f"/api/v1/squads/{squad_id}/matches",
            json={
                "seasonId": season_id,
                "date": f"2026-07-{index:02d}T14:00",
                "homeClub": home_club,
                "awayClub": away_club,
                "homeScore": home_score,
                "awayScore": away_score,
                "playerIds": player_ids,
            },
            headers={"X-CSRF-Token": token},
        )
        assert response.status_code == 201
        match_id = response.get_json()["match"]["id"]
        with app.app_context():
            charges = get_db().query(FineCharge).filter(FineCharge.match_id == match_id).order_by(FineCharge.player_id).all()
            assert [(charge.player_id, charge.amount_cents) for charge in charges] == [
                (player_ids[0], expected_cents),
                (player_ids[1], expected_cents),
            ]


def test_future_dbu_matches_are_visible_and_sorted(tmp_path, monkeypatch):
    from datetime import datetime

    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json={**setup_payload(), "dbuClubName": "Club"})
    squad_id = setup.get_json()["squad"]["id"]

    class FutureDBUClient:
        def __init__(self, **_kwargs):
            pass

        def list_season_match_ids(self, _season_url):
            return ["late", "early"]

        def fetch_match(self, match_id, _club_name):
            return {
                "dbu_id": match_id,
                "home_club": "CLUB",
                "away_club": match_id.upper(),
                "home_score": None,
                "away_score": None,
                "match_date": datetime(2026, 9 if match_id == "late" else 8, 1, 19, 0),
                "lineup": [],
            }

    monkeypatch.setattr("backend.services.DBUClient", FutureDBUClient)
    with app.app_context():
        database_session = get_db()
        report = sync_squad(database_session, database_session.get(Squad, squad_id))
        database_session.commit()
    assert report["createdMatches"] == 2
    matches = owner.get(f"/api/v1/squads/{squad_id}/dashboard").get_json()["matches"]
    assert [item["dbuId"] for item in matches] == ["early", "late"]
    assert all(item["homeScore"] is None and item["awayScore"] is None for item in matches)


def test_manual_match_and_player_deletion_rules(tmp_path):
    app = make_app(tmp_path)
    owner = app.test_client()
    setup = owner.post("/api/v1/setup", json=setup_payload())
    token = csrf(setup)
    squad_id = setup.get_json()["squad"]["id"]
    player = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Friendly Player"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    disposable = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Delete Me"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    assert owner.delete(
        f"/api/v1/squads/{squad_id}/players/{disposable['id']}",
        headers={"X-CSRF-Token": token},
    ).status_code == 200

    owner.post(
        f"/api/v1/squads/{squad_id}/fine-rules",
        json={"name": "Sejr", "type": "WIN_FINE", "amount": 10},
        headers={"X-CSRF-Token": token},
    )
    manual = owner.post(
        f"/api/v1/squads/{squad_id}/matches",
        json={
            "seasonId": setup.get_json()["squad"]["currentSeason"]["id"],
            "date": "2026-07-12T14:00",
            "homeClub": "Øster Sundby Boldklub",
            "awayClub": "Friends FC",
            "homeScore": 2,
            "awayScore": 1,
        },
        headers={"X-CSRF-Token": token},
    )
    assert manual.status_code == 201
    match_id = manual.get_json()["match"]["id"]
    assert manual.get_json()["match"]["dbuId"] is None
    assert owner.put(
        f"/api/v1/squads/{squad_id}/matches/{match_id}/lineup",
        json={"playerIds": [player["id"]]},
        headers={"X-CSRF-Token": token},
    ).status_code == 200
    with app.app_context():
        assert get_db().query(FineCharge).filter(FineCharge.match_id == match_id).count() == 1
    removed_match_player = owner.delete(
        f"/api/v1/squads/{squad_id}/players/{player['id']}",
        headers={"X-CSRF-Token": token},
    )
    assert removed_match_player.status_code == 200
    assert removed_match_player.get_json()["removedMatchFines"] == 1
    assert owner.delete(
        f"/api/v1/squads/{squad_id}/matches/{match_id}",
        headers={"X-CSRF-Token": token},
    ).status_code == 200
    with app.app_context():
        assert get_db().query(FineCharge).filter(FineCharge.match_id == match_id).count() == 0

    protected = owner.post(
        f"/api/v1/squads/{squad_id}/players",
        json={"dbuName": "Protected Player"},
        headers={"X-CSRF-Token": token},
    ).get_json()["player"]
    custom = owner.post(
        f"/api/v1/squads/{squad_id}/fine-requests",
        json={"title": "Bevar historik", "amount": 5, "playerIds": [protected["id"]]},
        headers={"X-CSRF-Token": token},
    )
    assert custom.status_code == 201
    member = app.test_client()
    assert member.post(
        "/api/v1/auth/register",
        json={"username": "leaving-player", "password": "", "playerId": protected["id"]},
    ).status_code == 201
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Name", "Type", "Number", "Message", "Amount", "Currency", "Transaction type"])
    sheet.append(["05/06/2026 20:21", "Protected Player", "Shop", "456", "", 10, "DKK", "Pay in"])
    file_data = BytesIO()
    workbook.save(file_data)
    imported = owner.post(
        f"/api/v1/squads/{squad_id}/mobilepay-imports",
        data={"file": (BytesIO(file_data.getvalue()), "protected-player.xlsx")},
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": token},
    )
    assert imported.status_code == 201

    removed = owner.delete(
        f"/api/v1/squads/{squad_id}/players/{protected['id']}",
        headers={"X-CSRF-Token": token},
    )
    assert removed.status_code == 200
    assert removed.get_json()["deletedAccount"] is True
    assert removed.get_json()["removedFines"] == 1
    assert removed.get_json()["removedFineRequests"] == 1
    assert removed.get_json()["unassignedPayments"] == 1
    assert member.get("/api/v1/auth/me").status_code == 401
    dashboard = app.test_client().get("/api/v1/public/squads/serie-1/dashboard").get_json()
    assert all(item["id"] != protected["id"] for item in dashboard["players"])
    assert dashboard["balanceSummary"]["boxBalance"] == 10
    transactions = owner.get(f"/api/v1/squads/{squad_id}/transactions").get_json()["transactions"]
    payment = next(item for item in transactions if item["name"] == "Protected Player")
    assert payment["allocatedPlayerId"] is None
    assert payment["allocationStatus"] == "unmatched"
    fine_requests = owner.get(f"/api/v1/squads/{squad_id}/fine-requests").get_json()["requests"]
    assert all(item["id"] != custom_request_id for item in fine_requests)
    with app.app_context():
        assert get_db().get(Player, protected["id"]) is None
        assert get_db().query(FineCharge).filter(FineCharge.player_id == protected["id"]).count() == 0
